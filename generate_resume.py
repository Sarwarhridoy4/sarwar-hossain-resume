#!/usr/bin/env python3
"""Generate Sarwar Hossain's ATS-friendly resume as a DOCX file."""

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


ACCENT = RGBColor(0x1F, 0x3A, 0x5F)   # dark navy for headings
SUBTLE = RGBColor(0x55, 0x55, 0x55)   # gray for contact/subtext
BODY = RGBColor(0x1A, 0x1A, 0x1A)     # near-black body text


def add_bottom_border(paragraph, size="14", color="1F3A5F"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    pBdr.append(bottom)
    pPr.append(pBdr)


def style_run(run, size=10.5, bold=False, italic=False, color=BODY, font="Lato"):
    run.font.name = font
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), font)
    rFonts.set(qn("w:hAnsi"), font)
    rFonts.set(qn("w:cs"), font)


def add_heading(doc, text, size=12, space_before=6, space_after=3):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.keep_with_next = True
    add_bottom_border(p, size="12", color="1F3A5F")
    run = p.add_run(text.upper())
    style_run(run, size=size, bold=True, color=ACCENT)
    return p


def add_body_paragraph(doc, text="", size=10, bold=False, italic=False,
                       color=BODY, space_after=3, space_before=0, align=None,
                       line_spacing=1.0):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.line_spacing = line_spacing
    if align is not None:
        p.alignment = align
    if text:
        run = p.add_run(text)
        style_run(run, size=size, bold=bold, italic=italic, color=color)
    return p


def add_bullet(doc, text, bold_prefix=None, size=10, space_after=2,
               indent=0.25, hang=0.15):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent)
    p.paragraph_format.first_line_indent = Inches(-hang)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.05
    marker = p.add_run("-  ")
    style_run(marker, size=size, bold=True, color=ACCENT)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        style_run(r1, size=size, bold=True, color=BODY)
    r2 = p.add_run(text)
    style_run(r2, size=size, color=BODY)
    return p


def add_experience_block(doc, employer, position, employment_type, dates,
                         location, bullets):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(f"{employer}  -  {position}")
    style_run(r, size=11, bold=True, color=ACCENT)

    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(4)
    p2.paragraph_format.keep_with_next = True
    r2 = p2.add_run(f"{employment_type}  -  {dates}  -  {location}")
    style_run(r2, size=10, italic=True, color=SUBTLE)

    for bp, text in bullets:
        add_bullet(doc, text, bold_prefix=bp)


def main():
    doc = Document()

    # Base style
    style = doc.styles["Normal"]
    style.font.name = "Lato"
    style.font.size = Pt(10.5)

    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # ---------- Name & Contact ----------
    name_p = doc.add_paragraph()
    name_p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    name_p.paragraph_format.space_after = Pt(1)
    add_bottom_border(name_p, size="8", color="1F3A5F")
    r = name_p.add_run("SARWAR HOSSAIN")
    style_run(r, size=20, bold=True, color=ACCENT)

    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(3)
    r = sub_p.add_run("Full Stack Software Engineer")
    style_run(r, size=12, bold=True, color=SUBTLE)

    # Contact line 1: location, phone, email
    contact_p1 = doc.add_paragraph()
    contact_p1.paragraph_format.space_after = Pt(1)
    first = True
    line1_items = ["Dhaka, Bangladesh", "+8801932893580", "sarwarhossainhridoy98@gmail.com"]
    for text in line1_items:
        if not first:
            sep = contact_p1.add_run("  -  ")
            style_run(sep, size=9.5, color=SUBTLE)
        run = contact_p1.add_run(text)
        style_run(run, size=9.5, color=SUBTLE)
        first = False

    # Contact line 2: portfolio, github, linkedin
    contact_p2 = doc.add_paragraph()
    contact_p2.paragraph_format.space_after = Pt(4)
    first = True
    line2_items = ["https://github.com/Sarwarhridoy4",
                   "https://sarwar.himulingua.com/",
                   "https://www.linkedin.com/in/sarwar-hossain-hridoy/"]
    for text in line2_items:
        if not first:
            sep = contact_p2.add_run("  -  ")
            style_run(sep, size=9.5, color=SUBTLE)
        run = contact_p2.add_run(text)
        style_run(run, size=9.5, color=SUBTLE)
        first = False

    # ---------- Professional Summary ----------
    add_heading(doc, "Professional Summary")
    add_body_paragraph(
        doc,
        "Full Stack Software Engineer with professional development experience "
        "since 2023, building business-facing web applications and a learning "
        "management system across the full stack. Practical experience with "
        "React, Next.js, TypeScript, Node.js, Express, MongoDB, and PostgreSQL, "
        "plus Go-based developer tools and automation. Skilled in deploying and "
        "troubleshooting applications on Linux VPS environments using Docker, "
        "reverse proxies, and SSL/TLS. Currently developing and maintaining the "
        "Himu Lingua LMS on a Hostinger VPS.",
        size=10,
        space_after=3,
    )

    # ---------- Technical Skills ----------
    add_heading(doc, "Technical Skills")

    skills_categories = [
        ("Frontend", "JavaScript, TypeScript, React.js, Next.js, HTML, CSS, Tailwind CSS, Redux Toolkit, responsive UI, reusable component architecture"),
        ("Backend", "Node.js, Express.js, Go, REST API development, authentication & authorization, JWT, bcrypt, MVC architecture, error handling, validation"),
        ("Databases", "MongoDB, PostgreSQL, Prisma ORM, Mongoose, database modeling, Zod schema validation, aggregation pipelines"),
        ("DevOps & Infrastructure", "Linux/Ubuntu, Docker, Hostinger VPS administration, Coolify, Traefik/Caddy, DNS, SSL/TLS, reverse proxies, application deployment, backup automation, GitHub Actions CI/CD"),
        ("Developer Tools", "Git, GitHub, shell scripting, API testing & debugging (Postman, Swagger), technical documentation, build & deployment troubleshooting"),
    ]

    for label, value in skills_categories:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1.0
        r0 = p.add_run(label + " - ")
        style_run(r0, size=10, bold=True, color=ACCENT)
        r1 = p.add_run(value)
        style_run(r1, size=10, color=BODY)

    # ---------- Professional Experience ----------
    add_heading(doc, "Professional Experience")

    add_experience_block(
        doc,
        employer="Himu Lingua (formerly German Ostad)",
        position="Full Stack Developer",
        employment_type="Remote, Project-Based",
        dates="August 2024 - Present",
        location="Remote",
        bullets=[
            ("Designed and developed a Learning Management System (LMS) from scratch",
             " using the MERN stack with TypeScript, including course-viewing interfaces, administrative dashboards, and custom authentication workflows."),
            ("Built REST APIs and backend services",
             " with Node.js and Express, integrating MongoDB data models, protected routes, session management, and video-based learning capabilities."),
            ("Deployed and maintained the application on a Hostinger VPS",
             " using Docker, Linux server administration, reverse proxy configuration, and SSL/TLS certificates; contributed to performance monitoring, reliability, and scalability."),
            ("Collaborated across frontend and backend development,",
             " application integration, server deployment, and ongoing technical improvements under a project-based engagement."),
        ],
    )

    add_experience_block(
        doc,
        employer="3W Private Limited",
        position="Full Stack Developer",
        employment_type="Full-time",
        dates="August 2023 - August 2024",
        location="India",
        bullets=[
            ("Delivered business-facing web applications",
             " using React.js and Express.js, spanning frontend UI development and backend API implementation."),
            ("Built backend API endpoints and MongoDB data models",
             " to support application functionality, including database-driven features and end-to-end web application development."),
            ("Developed dynamic, responsive user interfaces",
             " based on client requirements, emphasizing integration between frontend components and backend services."),
        ],
    )

    add_experience_block(
        doc,
        employer="Scanskill Edtech India Pvt. Ltd.",
        position="Full Stack Web Developer Intern",
        employment_type="Internship",
        dates="April 2023 - June 2023",
        location="India",
        bullets=[
            ("Fixed frontend and backend bugs",
             " and contributed to feature development on an existing web application."),
            ("Improved UI consistency and maintainability",
             " through reusable React components."),
            ("Contributed to API integrations and real-time application updates,",
             " helping improve application functionality and user experience."),
        ],
    )

    # ---------- Selected Projects ----------
    add_heading(doc, "Selected Projects")

    projects = [
        (
            "Digital Wallet Management System",
            "Express.js, TypeScript, MongoDB, Mongoose, JWT, Cloudinary",
            "Secure, role-based backend API modeled on digital wallet platforms. Implemented User, Agent, and Admin roles with isolated dashboards; wallet top-up, send money, withdraw, and cash-in/out flows with commission handling; JWT access/refresh tokens, email password reset, KYC verification, and account blocking.",
            "https://digital-wallet-management-system-kappa.vercel.app",
            "https://github.com/Sarwarhridoy4/Digital_Wallet_Management_System",
        ),
        (
            "Mongo Drive Backup",
            "Go, MongoDB, Google Drive API, Docker, Coolify",
            "Production-ready Go service that periodically backs up MongoDB databases and uploads them to Google Drive. Scheduled backups with cron, compression, retention cleanup, and upload verification; built-in web dashboard with real-time WebSocket updates for status, logs, and backup history; OAuth 2.0 and service-account support; Docker-ready with healthchecks and Coolify deployment.",
            None,
            "https://github.com/Sarwarhridoy4/mongo-drive-backup",
        ),
        (
            "Image Optimiser - Upload & Compression API",
            "Node.js, Express, TypeScript, MongoDB, Cloudinary, Sharp, Swagger, Nodemailer",
            "Secure image upload API with JWT authentication and access/refresh tokens. Compresses images to WebP using Sharp; Cloudinary integration with Multer upload and Zod validation; Swagger API documentation; welcome emails via Nodemailer with EJS templates.",
            "https://image-compress-upload.vercel.app",
            None,
        ),
        (
            "Express + Prisma MVC Boilerplate",
            "Express.js, TypeScript, Prisma, PostgreSQL, JWT, bcrypt",
            "Clean modular MVC boilerplate for scalable Express servers. Prisma ORM with PostgreSQL, JWT authentication, bcrypt password hashing, centralized error handling, and environment-based configuration; supports npm, yarn, pnpm, and bun.",
            None,
            "https://github.com/Sarwarhridoy4/demo-prisma-mvc-express",
        ),
        (
            "MongoDB Exporter GUI",
            "Python, PySide6, MongoDB, Cryptography, Google Drive API",
            "Desktop application that exports MongoDB collections to JSON with a neumorphic UI and theme switcher. Optional Zip compression, AES-GCM encryption with in-app decryption, Google Drive upload, backup scripts, real-time progress tracking, and built-in updater; cross-platform packaging (.deb and AppImage).",
            None,
            "https://github.com/Sarwarhridoy4/MongoDB-Exporter",
        ),
    ]

    for name, tech, desc, live, src in projects:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(5)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.keep_with_next = True
        r = p.add_run(name)
        style_run(r, size=11, bold=True, color=ACCENT)

        p2 = doc.add_paragraph()
        p2.paragraph_format.space_after = Pt(1)
        p2.paragraph_format.keep_with_next = True
        r2 = p2.add_run(tech)
        style_run(r2, size=10, italic=True, color=SUBTLE)

        add_body_paragraph(doc, desc, size=10, space_after=1)

        links = []
        if live:
            links.append(f"Live: {live}")
        if src:
            links.append(f"Source: {src}")
        if links:
            p3 = doc.add_paragraph()
            p3.paragraph_format.space_after = Pt(1)
            first = True
            for link in links:
                if not first:
                    sep = p3.add_run("   ")
                    style_run(sep, size=9, color=SUBTLE)
                run = p3.add_run(link)
                style_run(run, size=9, color=SUBTLE)
                first = False

    # ---------- Education ----------
    add_heading(doc, "Education")
    add_body_paragraph(
        doc,
        "Diploma in Engineering - SM Tofazzal Hossain Polytechnic Institute, 2021 (CGPA 3.67/4.0)",
        size=10,
        space_after=3,
    )

    # ---------- Certifications ----------
    add_heading(doc, "Certifications")
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run("Generative AI Mastermind - Outskill / GrowthSchool")
    style_run(r, size=10, bold=True, color=BODY)
    p2 = doc.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run("Certificate of Completion - Sarwar Hossain (issuer: Vaibhav Sisinity, Founder, GrowthSchool)")
    style_run(r2, size=10, italic=True, color=SUBTLE)

    p3 = doc.add_paragraph()
    p3.paragraph_format.space_after = Pt(1)
    r3 = p3.add_run("Complete Web Development - Programming Hero")
    style_run(r3, size=10, bold=True, color=BODY)
    p4 = doc.add_paragraph()
    p4.paragraph_format.space_after = Pt(0)
    r4 = p4.add_run("Certificate of Completion - Sarwar Hossain")
    style_run(r4, size=10, italic=True, color=SUBTLE)

    out_path = "/home/sarwar-hossain/Desktop/updated_resume/Sarwar_Hossain_Resume.docx"
    doc.save(out_path)
    print(f"Saved DOCX to {out_path}")


if __name__ == "__main__":
    main()
